"""Rebuild the DATA / TREND / MONTHS blocks in index.html from an IAB monthly report xlsx.

    python3 tools/build_data.py Podcasty_report_jul_2026.xlsx [--dry]

Add each new month to ALL_SHEETS (oldest first) as (sheet name, chart label); the
trend chart uses the last 12 entries and the newest one supplies the play counts.
Prints a report of name matches, stitched series and segment sizes - read it, the
source sheets rename podcasts between months.
"""
import openpyxl, json, re, unicodedata, difflib, sys, os

XLSX = next((a for a in sys.argv[1:] if a.endswith('.xlsx')), None)
HTML = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'index.html')
DRY = '--dry' in sys.argv
if not XLSX:
    sys.exit('usage: python3 tools/build_data.py <Podcasty_report_<month>.xlsx> [--dry]')

# monthly sheets, oldest -> newest; the last 12 form the trend window
ALL_SHEETS = [
    ('Máj 2025', 'Máj 25'), ('Jún 2025', 'Jún 25'), ('Júl 2025', 'Júl 25'),
    ('August 2025', 'Aug 25'), ('September 2025', 'Sep 25'), ('Október 2025', 'Okt 25'),
    ('November 2025', 'Nov 25'), ('December 2025', 'Dec 25'), ('Január_2026', 'Jan 26'),
    ('Februar_2026', 'Feb 26'), ('Marec _2026', 'Mar 26'), ('April_2026', 'Apr 26'),
    ('Máj_2026', 'Máj 26'), ('Jún_2026', 'Jún 26'), ('Júl_2026', 'Júl 26'),
]
WINDOW = ALL_SHEETS[-12:]
LATEST = ALL_SHEETS[-1][0]

# age-band midpoints: 0-17, 18-22, 23-27, 28-34, 35-44, 45-59, 60+
MID = [8.5, 20.0, 25.0, 31.0, 39.5, 52.0, 67.0]

# segment identities stay stable across updates; k-means is seeded from the previous
# centroids so cluster ids, labels and colours keep their meaning month to month
SEGMENTS = [
    dict(id=0, name='Muži 35+', sub='spravodajstvo, ekonomika, politika', color='#1971C2', seed=(28.0, 24.0)),
    dict(id=1, name='Mladšie ženy', sub='lifestyle, true crime, influenceri', color='#E64980', seed=(79.0, 56.0)),
    dict(id=2, name='Mladší muži', sub='šport, autá, technológie', color='#0CA678', seed=(14.0, 60.0)),
    dict(id=3, name='Ženy 35+', sub='vzťahy, zdravie, rodičovstvo', color='#7048E8', seed=(80.0, 28.0)),
    dict(id=4, name='Najmladšie publikum', sub='zmiešané, do 27 rokov', color='#F59F00', seed=(44.0, 71.0)),
    dict(id=5, name='Vyrovnané, stredný vek', sub='mix tém, 30 – 44 rokov', color='#495057', seed=(43.0, 45.0)),
]

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)


def key(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '', s.lower())


def parse_cell(v):
    """-> (value, already_in_percent)"""
    if v is None:
        return 0.0, True
    if isinstance(v, str):
        s = v.replace('%', '').replace('\xa0', '').replace(',', '.').strip()
        return (float(s), True) if s else (0.0, True)
    return float(v), False


def parse_group(vals):
    """Rows arrive either as fractions (0.325), as percentages (32.5) or as '32,5 %'."""
    parsed = [parse_cell(v) for v in vals]
    scale = 100.0 if sum(v for v, _ in parsed) <= 1.5 else 1.0
    return [round(v * (1.0 if pc else scale), 1) for v, pc in parsed]


def num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0
    try:
        return int(round(float(str(v).replace(',', '.'))))
    except (TypeError, ValueError):
        return 0


def month_rows(sheet):
    out = []
    for r in wb[sheet].iter_rows(min_row=3, values_only=True):
        if r[3] is None or str(r[3]).strip() == '':
            continue
        out.append(r)
    return out


# ---------------------------------------------------------------- podcasts (demografia)
podcasts, excluded = [], []
for r in wb['Demografia_2026'].iter_rows(min_row=3, values_only=True):
    if not r[2]:
        continue
    name = str(r[2]).strip()
    w, m, na, nb = parse_group(r[3:7])
    age = parse_group(r[7:14])
    if w + m + na + nb == 0 or sum(age) == 0:
        excluded.append(name)
        continue
    nz = sum(1 for x in age if x > 0)
    podcasts.append(dict(
        name=name,
        pub=str(r[0]).strip() if r[0] else '',
        med=str(r[1]).strip() if r[1] else '',
        cluster=0,
        w=w, m=m, na=na, nb=nb,
        age=age,
        wf=round(w / (w + m) * 100, 1) if (w + m) else 0.0,
        u35=round(sum(age[:4]), 1),
        meanage=round(sum(a * mid for a, mid in zip(age, MID)) / sum(age), 1),
        small=nz <= 2 or max(age) >= 80,
        plays=None,
        cat=None,
    ))

# ---------------------------------------------------------------- name matching
# The demografia sheet and the monthly sheets spell several podcasts differently
# ("ZKH" vs "Rozhovory ZKH", "Karieris" vs "Kariéééris"), so fall back from exact
# to normalised to a fuzzy match constrained to the same publisher.
def pub_ok(a, b):
    """"News and Media" (demografia) vs "News and Media Holding" (monthly sheets)."""
    ka, kb = key(a), key(b)
    return bool(ka) and bool(kb) and (ka == kb or ka.startswith(kb) or kb.startswith(ka))


def ratio(a, b):
    return difflib.SequenceMatcher(None, a, b).ratio()


def candidates_for(name, pub, candidates):
    """All candidate names that plausibly denote this podcast, best first."""
    kn = key(name)
    out = []
    for cn, cp in candidates:
        k = key(cn)
        if cn == name:
            score = 3.0
        elif k == kn:
            score = 2.0
        elif pub_ok(pub, cp) and min(len(k), len(kn)) >= 3 and (k in kn or kn in k):
            score = 1.0 + ratio(kn, k)
        elif (pub_ok(pub, cp) and ratio(kn, k) >= 0.75 and k[:4] == kn[:4]):
            # the shared prefix keeps near-anagrams apart ("kariéris" vs "kuriéris")
            score = ratio(kn, k)
        else:
            continue
        out.append((score, cn))
    out.sort(key=lambda t: -t[0])
    return [cn for _, cn in out]


def match(name, pub, candidates):
    c = candidates_for(name, pub, candidates)
    return c[0] if c else None


# ------------------------------------------------- plays + category (latest month)
latest = month_rows(LATEST)
lat_cands = [(str(r[3]).strip(), str(r[1] or '').strip()) for r in latest]
lat_by_name = {}
for r in latest:
    lat_by_name.setdefault(str(r[3]).strip(), r)

fuzzy_plays = []
for p in podcasts:
    hit = match(p['name'], p['pub'], lat_cands)
    if hit is None:
        continue
    if hit != p['name']:
        fuzzy_plays.append((p['name'], hit))
    r = lat_by_name[hit]
    total = r[12]
    p['plays'] = num(total) if num(total) else num(r[14]) + num(r[17]) + num(r[19])
    cat = r[5]
    p['cat'] = str(cat).strip() if cat and str(cat).strip() not in ('-', '') else None

# ---------------------------------------------------------------- clustering
def kmeans(points, seeds, iters=300):
    n = len(points)
    mx = sum(p[0] for p in points) / n
    my = sum(p[1] for p in points) / n
    sx = (sum((p[0] - mx) ** 2 for p in points) / n) ** .5 or 1.0
    sy = (sum((p[1] - my) ** 2 for p in points) / n) ** .5 or 1.0
    z = [((p[0] - mx) / sx, (p[1] - my) / sy) for p in points]
    cent = [((s[0] - mx) / sx, (s[1] - my) / sy) for s in seeds]
    assign = [-1] * n
    for _ in range(iters):
        new = [min(range(len(cent)), key=lambda k: (p[0] - cent[k][0]) ** 2 + (p[1] - cent[k][1]) ** 2)
               for p in z]
        if new == assign:
            break
        assign = new
        for k in range(len(cent)):
            mem = [z[i] for i in range(n) if assign[i] == k]
            if mem:
                cent[k] = (sum(q[0] for q in mem) / len(mem), sum(q[1] for q in mem) / len(mem))
    return assign


assign = kmeans([(p['wf'], p['u35']) for p in podcasts], [s['seed'] for s in SEGMENTS])
for p, a in zip(podcasts, assign):
    p['cluster'] = a

clusters = []
for s in SEGMENTS:
    mem = [p for p in podcasts if p['cluster'] == s['id']]
    n = len(mem)
    known = [p['plays'] for p in mem if p['plays'] is not None]
    clusters.append(dict(
        id=s['id'], name=s['name'], sub=s['sub'], color=s['color'], n=n,
        wf=round(sum(p['wf'] for p in mem) / n) if n else 0,
        meanage=round(sum(p['meanage'] for p in mem) / n) if n else 0,
        u35=round(sum(p['u35'] for p in mem) / n) if n else 0,
        age=[round(sum(p['age'][k] for p in mem) / n, 1) for k in range(7)] if n else [0] * 7,
        w=round(sum(p['w'] for p in mem) / n) if n else 0,
        m=round(sum(p['m'] for p in mem) / n) if n else 0,
        plays_total=sum(known), plays_known=len(known), plays_n=n,
    ))

known = [p['plays'] for p in podcasts if p['plays'] is not None]
DATA = dict(
    podcasts=podcasts,
    clusters=clusters,
    publishers=sorted(set(p['pub'] for p in podcasts)),
    biggest=max((p for p in podcasts if p['plays'] is not None), key=lambda p: p['plays'])['name'],
    plays_max=max(known),
    plays_min=min(known),
)

# ---------------------------------------------------------------- trend
months = [lab for _, lab in WINDOW]
N = len(WINDOW)
series, spub = {}, {}
for idx, (sheet, lab) in enumerate(WINDOW):
    for r in month_rows(sheet):
        nm = str(r[3]).strip()
        s = series.setdefault(nm, dict(a=[None] * N, y=[None] * N))
        spub[nm] = str(r[1] or '').strip()
        if s['a'][idx] is None:
            s['a'][idx] = num(r[14]) + num(r[17])
            s['y'][idx] = num(r[19])

# The monthly sheets re-spell and rename podcasts over time ("Zoom" -> "ZOOM",
# "Sketch Bros" -> "Ranná show Sketch Bros"). Stitch such series into one per
# podcast, but only where their reported months do not overlap - overlapping
# months would mean two distinct shows, so there we keep the best single series.
order = sorted(series, key=lambda k: -sum(v or 0 for v in series[k]['a'] + series[k]['y']))
cands = [(k, spub.get(k, '')) for k in order]
TREND, merges, no_trend = {}, [], []
for p in podcasts:
    names = candidates_for(p['name'], p['pub'], cands)
    if not names:
        no_trend.append(p['name'])
        continue
    used = [names[0]]
    for cn in names[1:]:
        if all(series[cn]['a'][i] is None or all(series[u]['a'][i] is None for u in used)
               for i in range(N)):
            used.append(cn)
    a = [next((series[u]['a'][i] for u in used if series[u]['a'][i] is not None), None) for i in range(N)]
    y = [next((series[u]['y'][i] for u in used if series[u]['y'][i] is not None), None) for i in range(N)]
    TREND[p['name']] = dict(a=a, y=y)
    if used != [p['name']]:
        merges.append((p['name'], used))

# ---------------------------------------------------------------- write
src = open(HTML).read()
lines = src.split('\n')
for i, ln in enumerate(lines):
    if ln.startswith('const DATA = '):
        lines[i] = 'const DATA = ' + json.dumps(DATA, ensure_ascii=False, separators=(', ', ': ')) + ';'
    elif ln.startswith('const TREND='):
        lines[i] = 'const TREND=' + json.dumps(TREND, ensure_ascii=False, separators=(',', ':')) + ';'
    elif ln.startswith('const MONTHS='):
        lines[i] = 'const MONTHS=' + json.dumps(months, ensure_ascii=False, separators=(',', ':')) + ';'
if not DRY:
    open(HTML, 'w').write('\n'.join(lines))

# ---------------------------------------------------------------- report
print('podcasts:', len(podcasts), '| excluded (no demographics):', excluded)
print('with plays:', len(known), '| without:', len(podcasts) - len(known))
print('   no plays ->', [p['name'] for p in podcasts if p['plays'] is None])
print('small sample:', [p['name'] for p in podcasts if p['small']])
print('fuzzy plays matches:')
for a, b in fuzzy_plays:
    print(f'   {a!r} -> {b!r}')
print('months:', months)
print('trend series:', len(TREND), 'of', len(series), 'raw | stitched:', len(merges))
for a, b in merges:
    print(f'   {a!r} <- {b}')
print('podcasts without trend:', no_trend)
print('biggest:', DATA['biggest'], DATA['plays_max'], '| min', DATA['plays_min'])
print('total plays (mapped):', sum(known))
for c in clusters:
    print(f"  {c['id']} {c['name']:24s} n={c['n']:3d} wf={c['wf']:3d} u35={c['u35']:3d} "
          f"age={c['meanage']} plays={c['plays_total']:>9d} known={c['plays_known']}")

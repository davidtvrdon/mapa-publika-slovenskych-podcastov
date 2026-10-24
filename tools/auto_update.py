#!/usr/bin/env python3
"""Roll the site forward when IAB Slovakia publishes a new monthly report.

Finds the newest report on https://www.iabslovakia.sk/category/data/, downloads the
xlsx linked from the article, registers the new month(s) in tools/build_data.py,
rebuilds the data and rewrites the period-dependent prose in index.html and README.md.

Exit codes: 0 = updated (or nothing to do), 1 = failed. Results for the workflow go to
$GITHUB_OUTPUT (status, label, branch, clean) and the PR body to tools/.pr_body.md.
`clean` is true when the rebuild shows nothing a human needs to look at (see is_clean).

  python3 tools/auto_update.py                       # detect, download, rebuild
  python3 tools/auto_update.py --xlsx F --article U  # skip the web lookup (testing)
"""
import html as htmllib, json, os, re, subprocess, sys, unicodedata, urllib.request
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUILD = os.path.join(ROOT, 'tools', 'build_data.py')
INDEX = os.path.join(ROOT, 'index.html')
README = os.path.join(ROOT, 'README.md')
BASELINE = os.path.join(ROOT, 'tools', 'report_baseline.json')
CATEGORY = 'https://www.iabslovakia.sk/category/data/'
UA = {'User-Agent': 'Mozilla/5.0 (mapa-podcastov auto-update)'}

# (normalised sheet/file prefix, short chart label, long Slovak name)
MONTHS = [('jan', 'Jan', 'január'), ('feb', 'Feb', 'február'), ('mar', 'Mar', 'marec'),
          ('apr', 'Apr', 'apríl'), ('maj', 'Máj', 'máj'), ('jun', 'Jún', 'jún'),
          ('jul', 'Júl', 'júl'), ('aug', 'Aug', 'august'), ('sep', 'Sep', 'september'),
          ('okt', 'Okt', 'október'), ('nov', 'Nov', 'november'), ('dec', 'Dec', 'december')]
LABEL_TO_NUM = {lab: i + 1 for i, (_, lab, _) in enumerate(MONTHS)}


def fail(msg):
    print('ERROR:', msg, file=sys.stderr)
    sys.exit(1)


def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '', s.lower())


def get(url, binary=False):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
        data = r.read()
    return data if binary else data.decode('utf-8', 'replace')


def month_num(word):
    n = norm(word)[:3]
    for i, (pre, _, _) in enumerate(MONTHS):
        if pre == n:
            return i + 1
    return None


def current_last_month():
    """(year, month) of the newest month already in index.html."""
    m = re.search(r'^const MONTHS=\[(.*?)\];', open(INDEX, encoding='utf-8').read(), re.M)
    last = json.loads('[' + m.group(1) + ']')[-1]
    lab, yy = last.split()
    return 2000 + int(yy), LABEL_TO_NUM[lab]


def find_latest_report():
    """-> (article_url, xlsx_url, (year, month)) for the newest report on the category page."""
    page = get(CATEGORY)
    links = []
    for u in re.findall(r'href="(https://www\.iabslovakia\.sk/[^"#?]+/)"', page):
        if 'vysledky-pocuvanosti-podcastov' in u and u not in links:
            links.append(u)
    best = None
    for u in links[:6]:
        for x in re.findall(r'href="([^"]*Podcasty_report_[^"]*\.xlsx)"', get(u), re.I):
            m = re.search(r'Podcasty_report_([a-z]+)_(\d{4})', x, re.I)
            mn = month_num(m.group(1)) if m else None
            if mn:
                ym = (int(m.group(2)), mn)
                if best is None or ym > best[2]:
                    best = (u, x, ym)
    if not best:
        fail('no report xlsx found on ' + CATEGORY)
    return best


def sheets_to_add(wb, last):
    """Sheets for every month after `last` up to the newest one the workbook holds."""
    found = {}
    for name in wb.sheetnames:
        ym = re.search(r'(20\d\d)', name)
        mn = month_num(name) if not name.startswith(('Demografia', 'Podcastov', 'Medziro')) else None
        if ym and mn:
            found.setdefault((int(ym.group(1)), mn), []).append(name)
    new = sorted(k for k in found if k > last)
    if not new:
        return []
    # months must follow each other, otherwise the trend chart would silently skip one
    y, m = last
    out = []
    for k in new:
        m += 1
        if m == 13:
            y, m = y + 1, 1
        if k != (y, m):
            fail(f'workbook skips a month: expected {y}-{m:02d}, next is {k[0]}-{k[1]:02d}')
        if len(found[k]) != 1:
            fail(f'ambiguous sheets for {k}: {found[k]}')
        out.append((found[k][0], f'{MONTHS[k[1] - 1][1]} {k[0] % 100}', k))
    return out


def sub1(pattern, repl, text, label, flags=0):
    new, n = re.subn(pattern, repl, text, count=1, flags=flags)
    if n != 1:
        fail(f'prose pattern not found, update tools/auto_update.py: {label}')
    return new


def best_excluded(wb, sheet, excluded):
    """Name and plays of the most-played podcast that has no demographics."""
    best = (None, 0)
    for r in wb[sheet].iter_rows(min_row=3, values_only=True):
        if r[3] and str(r[3]).strip() in excluded and isinstance(r[12], (int, float)) and r[12] > best[1]:
            best = (str(r[3]).strip(), r[12])
    return best


def is_clean(rep, base):
    """True when the rebuild holds nothing a human has to look at."""
    new_fuzzy = [t for t in rep['fuzzy'] if t not in base['fuzzy']]
    new_stitch = [t for t in rep['stitched'] if t not in base['stitched']]
    new_noplays = [n for n in rep['no_plays'] if n not in base['no_plays']]
    new_excl = [n for n in rep['excluded'] if n not in base['excluded']]
    return (rep['total'] == base['total'] and not new_fuzzy and not new_stitch
            and not new_excl and len(new_noplays) <= 3 and rep['months'][-1] != base['months'][-1])


def main():
    args = sys.argv[1:]
    opt = lambda k: args[args.index(k) + 1] if k in args else None
    last = current_last_month()
    if opt('--xlsx'):
        article, xlsx_path = opt('--article') or CATEGORY, opt('--xlsx')
    else:
        article, xlsx_url, ym = find_latest_report()
        print('newest report:', ym, xlsx_url)
        if ym <= last:
            print(f'already up to date ({last[0]}-{last[1]:02d})')
            out('status', 'uptodate')
            return
        xlsx_path = os.path.join(os.environ.get('RUNNER_TEMP', '/tmp'), 'report.xlsx')
        open(xlsx_path, 'wb').write(get(xlsx_url, binary=True))

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    add = sheets_to_add(wb, last)
    if not add:
        print('workbook holds no month after', last)
        out('status', 'uptodate')
        return

    # 1. register the new sheets
    src = open(BUILD, encoding='utf-8').read()
    entries = ''.join(f"    ('{sheet}', '{lab}'),\n" for sheet, lab, _ in add)
    src = sub1(r"(ALL_SHEETS = \[\n(?:.*\n)*?)(\]\n)", lambda m: m.group(1) + entries + m.group(2), src,
               'ALL_SHEETS', re.M)
    open(BUILD, 'w', encoding='utf-8').write(src)

    # 2. rebuild data
    rep_path = os.path.join(ROOT, 'tools', '.report.json')
    proc = subprocess.run([sys.executable, '-I', BUILD, xlsx_path, '--report-json', rep_path],
                          capture_output=True, text=True)
    sys.stdout.write(proc.stdout)
    if proc.returncode:
        fail('build_data.py failed:\n' + proc.stderr)
    rep = json.load(open(rep_path))
    base = json.load(open(BASELINE))
    clean = is_clean(rep, base)

    # 3. prose
    (y, mn) = add[-1][2]
    long_m = MONTHS[mn - 1][2]
    period = f'{long_m} {y}'
    first = rep['months'][0].split()
    start = f"{MONTHS[LABEL_TO_NUM[first[0]] - 1][2]} {2000 + int(first[1])}"
    today = datetime.now(ZoneInfo('Europe/Bratislava'))
    stamp = f'{today.day}. {today.month}. {today.year}'
    mmyyyy = f'{mn:02d}/{y}'
    big_name, big_plays = best_excluded(wb, add[-1][0], rep['excluded'])
    mil = f'{big_plays / 1e6:.1f}'.replace('.', ',')
    excl_txt = ', '.join(rep['excluded'])
    total, withp = rep['total'], rep['plays']

    h = open(INDEX, encoding='utf-8').read()
    h = sub1(r'(Dátový zdroj: <a href=")[^"]*(" target="_blank" rel="noopener noreferrer">IAB Slovakia: '
             r'Monitoring počúvanosti podcastov \()\d\d/\d{4}(\)</a>)( · Aktualizované [^<]*)?',
             lambda m: f'{m.group(1)}{article}{m.group(2)}{mmyyyy}{m.group(3)} · Aktualizované {stamp}', h, 'source note')
    h = sub1(r'Počty prehratí sú za \S+ \d{4} a podarilo sa ich priradiť k \d+ zo \d+ podcastov; \d+ podcastov',
             f'Počty prehratí sú za {period} a podarilo sa ich priradiť k {withp} zo {total} podcastov; '
             f'{total - withp} podcastov', h, 'methodology counts')
    h = sub1(r'(bez demografických dát \()[^,()]+, [\d,]+( mil\.\) na mape nie je; najväčší zobrazený bod je )[^.;<]+\.',
             lambda m: f'{m.group(1)}{big_name}, {mil}{m.group(2)}{rep["biggest"]}.', h, 'largest podcast')
    h = sub1(r'Podcasty bez demografie \([^)]*\) boli vynechané',
             f'Podcasty bez demografie ({htmllib.escape(excl_txt, quote=False)}) boli vynechané', h, 'excluded list')
    open(INDEX, 'w', encoding='utf-8').write(h)

    r = open(README, encoding='utf-8').read()
    r = sub1(r'\(\d\d/\d{4}\)\]\([^)]*\)( · aktualizované [^\n]*)?',
             f'({mmyyyy})]({article}) · aktualizované {stamp}', r, 'readme source')
    r = sub1(r'celkový počet prehratí za \S+ \d{4} \(', f'celkový počet prehratí za {period} (', r, 'readme size')
    r = sub1(r'\(\S+ \d{4} – \S+ \d{4}\)', f'({start} – {period})', r, 'readme range')
    r = sub1(r'Podcasty bez demografických dát \([^)]*\)', f'Podcasty bez demografických dát ({excl_txt})', r, 'readme excluded')
    r = sub1(r'z prehratí za \S+ \d{4}, ktoré sa podarilo priradiť k \d+ zo \d+ podcastov – zvyšných \d+ je',
             f'z prehratí za {period}, ktoré sa podarilo priradiť k {withp} zo {total} podcastov – zvyšných {total - withp} je', r, 'readme counts')
    r = sub1(r'teda \S+ \d{4} – \S+ \d{4}\.', f'teda {start} – {period}.', r, 'readme trend')
    open(README, 'w', encoding='utf-8').write(r)

    # 4. PR text + outputs
    new_fuzzy = [t for t in rep['fuzzy'] if t not in base['fuzzy']]
    new_stitch = [t for t in rep['stitched'] if t not in base['stitched']]
    new_noplays = [n for n in rep['no_plays'] if n not in base['no_plays']]
    body = [f'Automatic update to **{period}** from [{mmyyyy}]({article}).', '',
            f'- Podcasts: {total} ({withp} with plays, {total - withp} without)',
            f'- Largest without demographics: {big_name} ({mil} mil.)',
            f'- New fuzzy name matches: {new_fuzzy or "none"}',
            f'- New stitched trend series: {new_stitch or "none"}',
            f'- Newly missing play counts: {new_noplays or "none"}', '',
            ('Nothing new to review: this PR is auto-merged and published.' if clean else
             '**Needs a human look** (new matches, a changed podcast list or many missing plays), '
             'so it is not auto-merged. Check the fuzzy matches and merge by hand.'),
            '', 'Not refreshed automatically: `seed` centroids in `SEGMENTS` (see CLAUDE.md step 5).']
    open(os.path.join(ROOT, 'tools', '.pr_body.md'), 'w', encoding='utf-8').write('\n'.join(body) + '\n')
    json.dump(rep, open(BASELINE, 'w'), ensure_ascii=False, indent=1)
    os.remove(rep_path)
    out('status', 'updated')
    out('label', period)
    out('branch', f'data/{y}-{mn:02d}')
    out('clean', 'true' if clean else 'false')
    print('updated to', period, '| clean:', clean)


def out(k, v):
    print(f'{k}={v}')
    if os.environ.get('GITHUB_OUTPUT'):
        open(os.environ['GITHUB_OUTPUT'], 'a').write(f'{k}={v}\n')


if __name__ == '__main__':
    main()
